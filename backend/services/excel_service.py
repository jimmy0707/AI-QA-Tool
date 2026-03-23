"""
services/excel_service.py
─────────────────────────
Excel report generators for both Regression and Automation analysis.

Features:
  - AI Source column with colour-coded badges per engine
  - Release-aware columns auto-detected (Base/Adjusted score, Adjustment Reason)
  - Summary sheet with AI Source Breakdown + Release Adjustment tables
  - Full colour legend on Summary sheet
"""

import os
import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

from config import OUTPUT_DIR
from utils.helpers import safe_str


# ── AI Source colour palette ─────────────────────────────────────────────────
_AI_SOURCE_FILLS = {
    "gemini":          PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
    "gemini-fallback": PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"),
    "openai":          PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid"),
    "openai-fallback": PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"),
    "offline":         PatternFill(start_color="EDE7F6", end_color="EDE7F6", fill_type="solid"),
    "heuristic":       PatternFill(start_color="F3E5F5", end_color="F3E5F5", fill_type="solid"),
}
_AI_SOURCE_FONT_COLORS = {
    "gemini":          "1B5E20",
    "gemini-fallback": "F57F17",
    "openai":          "0D47A1",
    "openai-fallback": "F57F17",
    "offline":         "4A148C",
    "heuristic":       "6A1B9A",
}
_AI_SOURCE_LABELS = {
    "gemini":          "🟢 Google Gemini",
    "gemini-fallback": "⚠️ Gemini→Fallback",
    "openai":          "🔵 OpenAI GPT",
    "openai-fallback": "⚠️ OpenAI→Fallback",
    "offline":         "🔌 Local Ollama",
    "heuristic":       "📋 Keyword Rules",
}

# Adjustment reason fills
_ADJ_FILLS = {
    "changed": PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),
    "frozen":  PatternFill(start_color="BBDEFB", end_color="BBDEFB", fill_type="solid"),
    "stale":   PatternFill(start_color="FFE0B2", end_color="FFE0B2", fill_type="solid"),
    "none":    PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"),
}


# ── Helper functions ─────────────────────────────────────────────────────────

def _ai_source_label(raw_mode: str) -> str:
    key = str(raw_mode).lower().strip()
    return _AI_SOURCE_LABELS.get(key, f"🤖 {raw_mode.upper()}")


def _apply_ai_source_cell(cell, raw_mode: str):
    key    = str(raw_mode).lower().strip()
    fill   = _AI_SOURCE_FILLS.get(key, PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid"))
    fcolor = _AI_SOURCE_FONT_COLORS.get(key, "424242")
    cell.fill      = fill
    cell.font      = Font(bold=True, color=fcolor, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _engine_banner(mode: str) -> str:
    return {
        "gemini":  "Google Gemini AI",
        "openai":  "OpenAI GPT-4o-mini",
        "offline": "Local Ollama (Offline AI)",
    }.get(mode.lower(), mode.upper())


def _adj_fill(reason_str: str):
    r = (reason_str or "").lower()
    if "changed" in r: return _ADJ_FILLS["changed"]
    if "frozen"  in r: return _ADJ_FILLS["frozen"]
    if "stale"   in r: return _ADJ_FILLS["stale"]
    return _ADJ_FILLS["none"]


def _write_summary_engine_section(ws, mode: str):
    """Append AI Engine info + colour legend to a Summary sheet."""
    ws.append([])
    ws.append(["─── AI Engine Information ───", "", ""])
    ws.append(["Primary AI Engine", _engine_banner(mode), ""])
    ws.append(["Fallback Engine",   "Local Ollama → Keyword Rules", ""])
    ws.append(["Report Generated",  datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), ""])
    ws.append([])
    ws.append(["AI Source Legend", "", ""])

    legend = [
        ("🟢 Google Gemini",   "E8F5E9", "1B5E20", "Result produced by Google Gemini AI"),
        ("🔵 OpenAI GPT",      "E3F2FD", "0D47A1", "Result produced by OpenAI GPT-4o-mini"),
        ("🔌 Local Ollama",    "EDE7F6", "4A148C", "Result produced by local Ollama model"),
        ("📋 Keyword Rules",   "F3E5F5", "6A1B9A", "Result from keyword heuristic engine"),
        ("⚠️ Gemini→Fallback", "FFF9C4", "F57F17", "Gemini failed — result from fallback engine"),
        ("⚠️ OpenAI→Fallback", "FFF9C4", "F57F17", "OpenAI failed — result from fallback engine"),
    ]
    for label, bg, fg, desc in legend:
        r = ws.max_row + 1
        ws.append([label, desc, ""])
        ws.cell(row=r, column=1).fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
        ws.cell(row=r, column=1).font = Font(bold=True, color=fg, size=10)
        ws.cell(row=r, column=2).font = Font(color="424242", size=10)


def _write_ai_source_breakdown(ws, df: pd.DataFrame, total: int):
    """Append the AI Source Breakdown table to a Summary sheet."""
    ws.append([])
    ws.append(["AI Source Breakdown", "Count", "% of Total"])
    hdr_row  = ws.max_row
    src_fill = PatternFill(start_color="37474F", end_color="37474F", fill_type="solid")
    for col in range(1, 4):
        c = ws.cell(row=hdr_row, column=col)
        c.fill = src_fill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")

    source_col = "AI Source" if "AI Source" in df.columns else "AI Mode"
    if source_col in df.columns:
        for src_val, count in df[source_col].value_counts().items():
            pct   = f"{round(count / total * 100)}%" if total else "0%"
            label = _ai_source_label(str(src_val))
            r     = ws.max_row + 1
            ws.append([label, count, pct])
            _apply_ai_source_cell(ws.cell(row=r, column=1), str(src_val))
            ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
            ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")


# ── Public generators ────────────────────────────────────────────────────────

def create_regression_excel(df: pd.DataFrame, session_id: str, mode: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Regression Results"

    has_release = "Adjusted Risk Score" in df.columns

    # ── Headers ──────────────────────────────────────────────────────────────
    if has_release:
        headers = ["#", "Test Case Title", "Description", "Module",
                   "Base Risk Score", "Adjusted Risk Score", "Adjustment Reason",
                   "Priority", "AI Risk Explanation", "Recommended", "AI Source"]
    else:
        headers = ["#", "Test Case Title", "Description",
                   "Risk Score", "Priority", "AI Risk Explanation",
                   "Recommended", "AI Source"]

    ws.append(headers)
    hfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = hfill
        c.font = Font(color="FFFFFF", bold=True, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ── Data rows ─────────────────────────────────────────────────────────────
    priority_colors = {"P1": "FFCCCC", "P2": "FFF2CC", "P3": "D9EAD3"}

    for i, row in df.iterrows():
        title       = safe_str(row.get("Title") or row.get("title") or row.get("Test Case Title"), f"TC-{i+1}")
        desc        = safe_str(row.get("Description") or row.get("description"), "")
        priority    = str(row.get("Priority", "P2"))
        explanation = safe_str(row.get("AI Risk Explanation"), "")
        recommended = str(row.get("Recommended for Execution", "No"))
        raw_source  = str(row.get("AI Source") or row.get("AI Mode") or mode)
        src_label   = _ai_source_label(raw_source)
        data_row    = i + 2

        color = priority_colors.get(priority, "FFFFFF")
        rfill = PatternFill(start_color=color, end_color=color, fill_type="solid")

        if has_release:
            module     = safe_str(row.get("Module") or row.get("module"), "—")
            base_score = row.get("Base Risk Score", 5)
            adj_score  = row.get("Adjusted Risk Score", base_score)
            adj_reason = safe_str(row.get("Adjustment Reason"), "No adjustment")

            ws.append([i + 1, title, desc, module, base_score, adj_score,
                       adj_reason, priority, explanation, recommended, src_label])

            for col in list(range(1, 7)) + [8, 9, 10]:
                c = ws.cell(row=data_row, column=col)
                c.fill = rfill
                c.alignment = Alignment(wrap_text=True, vertical="top")

            reason_cell = ws.cell(row=data_row, column=7)
            reason_cell.fill = _adj_fill(adj_reason)
            reason_cell.font = Font(size=10, italic=True)
            reason_cell.alignment = Alignment(wrap_text=True, vertical="top")

            _apply_ai_source_cell(ws.cell(row=data_row, column=11), raw_source)

        else:
            risk_score = row.get("Risk Score", 5)
            ws.append([i + 1, title, desc, risk_score, priority, explanation, recommended, src_label])
            for col in range(1, 8):
                c = ws.cell(row=data_row, column=col)
                c.fill = rfill
                c.alignment = Alignment(wrap_text=True, vertical="top")
            _apply_ai_source_cell(ws.cell(row=data_row, column=8), raw_source)

    # ── Column widths ─────────────────────────────────────────────────────────
    widths = (
        {"A": 6, "B": 32, "C": 38, "D": 16, "E": 14, "F": 16, "G": 45, "H": 10, "I": 48, "J": 14, "K": 22}
        if has_release else
        {"A": 6, "B": 35, "C": 40, "D": 12, "E": 10, "F": 50, "G": 15, "H": 22}
    )
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.row_dimensions[1].height = 35

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2   = wb.create_sheet("Summary")
    p1    = len(df[df["Priority"] == "P1"])
    p2    = len(df[df["Priority"] == "P2"])
    p3    = len(df[df["Priority"] == "P3"])
    total = len(df)
    rec   = len(df[df["Recommended for Execution"] == "Yes"])

    ws2.append(["AI QA Decision Intelligence — Regression Report", "", ""])
    ws2.merge_cells("A1:C1")
    ws2["A1"].font      = Font(bold=True, size=14, color="1E3A5F")
    ws2["A1"].alignment = Alignment(horizontal="center")

    mode_line = f"AI Engine: {_engine_banner(mode)}"
    if has_release:
        mode_line += "  |  Release-Aware Prioritization: ON"
    ws2.append([mode_line, "", ""])
    ws2.merge_cells("A2:C2")
    ws2["A2"].font      = Font(bold=True, size=11, color="2E75B6")
    ws2["A2"].alignment = Alignment(horizontal="center")

    ws2.append([])
    ws2.append(["Metric", "Count", "Percentage"])
    ws2.append(["Total Test Cases",          total, "100%"])
    ws2.append(["P1 — Critical",             p1,    f"{round(p1/total*100)}%" if total else "0%"])
    ws2.append(["P2 — Moderate",             p2,    f"{round(p2/total*100)}%" if total else "0%"])
    ws2.append(["P3 — Low Risk",             p3,    f"{round(p3/total*100)}%" if total else "0%"])
    ws2.append(["Recommended for Execution", rec,   f"{round(rec/total*100)}%" if total else "0%"])

    sfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for col in range(1, 4):
        c = ws2.cell(row=4, column=col)
        c.fill = sfill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")

    # Release adjustment summary
    if has_release and "Adjustment Reason" in df.columns:
        ws2.append([])
        ws2.append(["Release Adjustment Summary", "Count", "% of Total"])
        adj_hdr = ws2.max_row
        adj_fill_hdr = PatternFill(start_color="37474F", end_color="37474F", fill_type="solid")
        for col in range(1, 4):
            c = ws2.cell(row=adj_hdr, column=col)
            c.fill = adj_fill_hdr
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center")

        changed_ct = df["Adjustment Reason"].str.contains("Changed", na=False).sum()
        frozen_ct  = df["Adjustment Reason"].str.contains("Frozen",  na=False).sum()
        stale_ct   = df["Adjustment Reason"].str.contains("Stale",   na=False).sum()
        none_ct    = total - changed_ct - frozen_ct - stale_ct

        for label, count, bg, fg in [
            ("🟢 Changed module (boosted)", changed_ct, "C8E6C9", "1B5E20"),
            ("🔵 Frozen module (reduced)",  frozen_ct,  "BBDEFB", "0D47A1"),
            ("🟠 Stale test (boosted)",      stale_ct,  "FFE0B2", "E65100"),
            ("⬜ No adjustment",             none_ct,   "F5F5F5", "424242"),
        ]:
            pct = f"{round(count/total*100)}%" if total else "0%"
            r   = ws2.max_row + 1
            ws2.append([label, count, pct])
            ws2.cell(row=r, column=1).fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            ws2.cell(row=r, column=1).font = Font(bold=True, color=fg, size=10)
            ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center")
            ws2.cell(row=r, column=3).alignment = Alignment(horizontal="center")

    _write_ai_source_breakdown(ws2, df, total)
    _write_summary_engine_section(ws2, mode)

    for col in ["A", "B", "C"]:
        ws2.column_dimensions[col].width = 36

    path = os.path.join(OUTPUT_DIR, f"regression_results_{session_id}.xlsx")
    wb.save(path)
    return path


def create_automation_excel(df: pd.DataFrame, session_id: str, mode: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Automation Analysis"

    headers = ["#", "Test Case Title", "Description",
               "Automation Suitability", "Confidence %", "AI Explanation", "AI Source"]
    ws.append(headers)

    hfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = hfill
        c.font = Font(color="FFFFFF", bold=True, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")

    suit_colors = {"Automatable": "D9EAD3", "Partial": "FFF2CC", "Not Suitable": "FFCCCC"}

    for i, row in df.iterrows():
        title       = safe_str(row.get("Title") or row.get("title") or row.get("Test Case Title"), f"TC-{i+1}")
        desc        = safe_str(row.get("Description") or row.get("description"), "")
        suitability = str(row.get("Automation Suitability", "Partial"))
        confidence  = row.get("Confidence %", 50)
        explanation = safe_str(row.get("AI Explanation"), "")
        raw_source  = str(row.get("AI Source") or row.get("AI Mode") or mode)
        src_label   = _ai_source_label(raw_source)
        data_row    = i + 2

        ws.append([i + 1, title, desc, suitability, f"{confidence}%", explanation, src_label])

        color = suit_colors.get(suitability, "FFFFFF")
        rfill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        for col in range(1, 7):
            c = ws.cell(row=data_row, column=col)
            c.fill = rfill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        _apply_ai_source_cell(ws.cell(row=data_row, column=7), raw_source)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 55
    ws.column_dimensions["G"].width = 22
    ws.row_dimensions[1].height = 35

    ws2          = wb.create_sheet("Summary")
    total        = len(df)
    automatable  = len(df[df["Automation Suitability"] == "Automatable"])
    partial      = len(df[df["Automation Suitability"] == "Partial"])
    not_suitable = len(df[df["Automation Suitability"] == "Not Suitable"])

    ws2.append(["AI QA Decision Intelligence — Automation Report", "", ""])
    ws2.merge_cells("A1:C1")
    ws2["A1"].font      = Font(bold=True, size=14, color="1E3A5F")
    ws2["A1"].alignment = Alignment(horizontal="center")

    ws2.append([f"AI Engine: {_engine_banner(mode)}", "", ""])
    ws2.merge_cells("A2:C2")
    ws2["A2"].font      = Font(bold=True, size=11, color="2E75B6")
    ws2["A2"].alignment = Alignment(horizontal="center")

    ws2.append([])
    ws2.append(["Category", "Count", "Percentage"])
    ws2.append(["Total Analyzed",  total,        "100%"])
    ws2.append(["Automatable",     automatable,  f"{round(automatable/total*100)}%"  if total else "0%"])
    ws2.append(["Partial",         partial,      f"{round(partial/total*100)}%"      if total else "0%"])
    ws2.append(["Not Suitable",    not_suitable, f"{round(not_suitable/total*100)}%" if total else "0%"])

    sfill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    for col in range(1, 4):
        c = ws2.cell(row=4, column=col)
        c.fill = sfill
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center")

    _write_ai_source_breakdown(ws2, df, total)
    _write_summary_engine_section(ws2, mode)

    for col in ["A", "B", "C"]:
        ws2.column_dimensions[col].width = 32

    path = os.path.join(OUTPUT_DIR, f"automation_results_{session_id}.xlsx")
    wb.save(path)
    return path
